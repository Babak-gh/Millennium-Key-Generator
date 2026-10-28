import datetime
import io
import os
import unittest
from unittest.mock import Mock, patch

from openpyxl import load_workbook
from app import app, db, User, License, ZigbeeLicense
from cloud_billing import reconcile_cloud_coverage, ReconciliationUnavailable, excel_text

TOKEN = 't' * 40


class CloudBillingTest(unittest.TestCase):
    def setUp(self):
        self.environment = patch.dict(os.environ, {
            'LICENSING_SERVICE_TOKEN': TOKEN, 'CLOUD_LICENSING_URL': 'https://cloud.example.test',
        })
        self.environment.start()
        app.config['TESTING'] = True
        self.client = app.test_client()
        with app.app_context():
            License.query.filter_by(code='cloud-billing-panel').delete()
            ZigbeeLicense.query.filter_by(code='cloud-billing-panel').delete()
            db.session.add(License(code='cloud-billing-panel', issuer='issuer', owner='=malicious()',
                                   project='Home', is_active=True, created_date=datetime.datetime.now(), license='base'))
            db.session.add(ZigbeeLicense(code='cloud-billing-panel', license_id='cloud-billing-zigbee',
                                        coordinator_eui64='00124B0000FFFFFF', issuer='issuer', owner='Owner',
                                        project='Home', is_active=True, created_date=datetime.datetime.now(), license='zigbee'))
            db.session.commit()
            self.base_key = 'BASE:' + str(License.query.filter_by(code='cloud-billing-panel').one().id)
            self.admin_id = str(User.query.filter_by(username='admin').one().id)
        with self.client.session_transaction() as session:
            session['_user_id'] = self.admin_id
            session['_fresh'] = True

    def tearDown(self):
        self.environment.stop()
        with app.app_context():
            License.query.filter_by(code='cloud-billing-panel').delete()
            ZigbeeLicense.query.filter_by(code='cloud-billing-panel').delete()
            db.session.commit()

    def test_lookup_requires_service_authentication_and_returns_no_signatures(self):
        url = '/internal/cloud/licenses?code=cloud-billing-panel'
        self.assertEqual(401, self.client.get(url).status_code)
        response = self.client.get(url, headers={'Authorization': 'Bearer ' + TOKEN})
        self.assertEqual(200, response.status_code)
        records = response.json['licenses']
        self.assertEqual({'BUS', 'ZIGBEE'}, {row['feature'] for row in records})
        self.assertTrue(all('license' not in row for row in records))

    @patch('cloud_billing.requests.post')
    def test_export_contains_both_features_but_excludes_cloud_covered_from_collection(self, post):
        def reply(*args, **kwargs):
            rows = []
            for row in kwargs['json']['licenses']:
                covered = row['feature'] == 'BUS'
                rows.append({'key': row['key'], 'status': 'COVERED_BY_CLOUD' if covered else 'REVIEW_REQUIRED',
                             'homeId': 'home' if covered else None, 'paymentId': 'payment' if covered else None,
                             'reference': 'transfer' if covered else None})
            return Mock(status_code=200, json=lambda: {'results': rows})
        post.side_effect = reply
        response = self.client.get('/admin/license/export_billing')
        self.assertEqual(200, response.status_code)
        workbook = load_workbook(io.BytesIO(response.data))
        rows = list(workbook['All licenses'].values)
        self.assertEqual(3, len(rows))
        self.assertEqual("'=malicious()", rows[1][4])
        self.assertEqual(1, workbook['Offline payment review'].max_row)
        self.assertEqual(2, workbook['Cloud review required'].max_row)
        self.assertFalse(post.call_args.kwargs['allow_redirects'])

    @patch('cloud_billing.requests.post')
    def test_failed_or_incomplete_reconciliation_never_produces_a_collection_export(self, post):
        for response in [Mock(status_code=503), Mock(status_code=200, json=lambda: {'results': []})]:
            post.return_value = response
            result = self.client.get('/admin/license/export_billing')
            self.assertEqual(503, result.status_code)
            self.assertIn('No collection report', result.json['error'])

    def test_unconfigured_connection_is_not_treated_as_unpaid(self):
        with patch.dict(os.environ, {'LICENSING_SERVICE_TOKEN': ''}):
            with self.assertRaises(ReconciliationUnavailable):
                reconcile_cloud_coverage([])

    @patch('cloud_billing.requests.post')
    def test_service_failure_on_a_later_batch_fails_entire_report(self, post):
        records = [{'key': str(i)} for i in range(501)]
        post.side_effect = [Mock(status_code=200, json=lambda: {'results': [
            {'key': str(i), 'status': 'NOT_COVERED'} for i in range(500)]}), Mock(status_code=503)]
        with self.assertRaises(ReconciliationUnavailable):
            reconcile_cloud_coverage(records)
        self.assertEqual(2, post.call_count)

    def test_formula_prefixes_are_escaped(self):
        self.assertEqual("'  =SUM(A1)", excel_text('  =SUM(A1)'))
        self.assertEqual('panel-123', excel_text('panel-123'))


if __name__ == '__main__':
    unittest.main()
